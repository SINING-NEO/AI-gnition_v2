"""Train the first purchase-timing and next-product models.

The input workbook contains synthetic data only. Splits are predefined at the
customer level, so rows from one customer never leak between train and test.

This module intentionally depends only on NumPy and openpyxl, which are already
available in the project environment.
"""

from __future__ import annotations

import argparse
import csv
import json
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DATA = ROOT / "data" / "retail_purchase_training_dataset.xlsx"
DEFAULT_OUTPUT = ROOT / "prediction" / "artifacts"

NUMERIC_FEATURES = [
    "order_count_to_date",
    "days_since_previous_order",
    "mean_gap_days",
    "median_gap_days",
    "gap_std_days",
    "dominant_product_share",
    "same_as_previous_product",
    "latest_quantity",
    "latest_order_value_sgd",
    "latest_discount_pct",
]

CATEGORICAL_FEATURES = [
    "dominant_product_id",
    "latest_product_id",
    "previous_product_id",
    "latest_channel",
    "region",
    "customer_profile",
]


def load_rows(path: Path) -> list[dict]:
    """Load the model-ready sheet without changing its stored values."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Model Training"]
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    rows = [dict(zip(headers, values)) for values in iterator]
    workbook.close()
    if not rows:
        raise ValueError("The Model Training sheet contains no data rows.")
    return rows


@dataclass
class FeatureEncoder:
    numeric_features: list[str]
    categorical_features: list[str]
    means: np.ndarray | None = None
    scales: np.ndarray | None = None
    categories: dict[str, list[str]] | None = None

    def fit(self, rows: list[dict]) -> "FeatureEncoder":
        numeric = self._numeric_matrix(rows)
        self.means = numeric.mean(axis=0)
        self.scales = numeric.std(axis=0)
        self.scales[self.scales == 0] = 1.0
        self.categories = {
            feature: sorted({str(row[feature]) for row in rows})
            for feature in self.categorical_features
        }
        return self

    def _numeric_matrix(self, rows: list[dict]) -> np.ndarray:
        return np.asarray(
            [
                [float(bool(row[name])) if isinstance(row[name], bool) else float(row[name])
                 for name in self.numeric_features]
                for row in rows
            ],
            dtype=np.float64,
        )

    def transform(self, rows: list[dict]) -> np.ndarray:
        if self.means is None or self.scales is None or self.categories is None:
            raise RuntimeError("FeatureEncoder must be fitted before transform().")
        blocks = [(self._numeric_matrix(rows) - self.means) / self.scales]
        for feature in self.categorical_features:
            category_to_index = {
                category: index for index, category in enumerate(self.categories[feature])
            }
            encoded = np.zeros((len(rows), len(category_to_index)), dtype=np.float64)
            for row_index, row in enumerate(rows):
                category_index = category_to_index.get(str(row[feature]))
                if category_index is not None:
                    encoded[row_index, category_index] = 1.0
            blocks.append(encoded)
        return np.column_stack(blocks)

    def feature_names(self) -> list[str]:
        if self.categories is None:
            raise RuntimeError("FeatureEncoder has not been fitted.")
        names = list(self.numeric_features)
        for feature in self.categorical_features:
            names.extend(f"{feature}={value}" for value in self.categories[feature])
        return names

    def metadata(self) -> dict:
        if self.means is None or self.scales is None or self.categories is None:
            raise RuntimeError("FeatureEncoder has not been fitted.")
        return {
            "numeric_features": self.numeric_features,
            "categorical_features": self.categorical_features,
            "means": self.means.tolist(),
            "scales": self.scales.tolist(),
            "categories": self.categories,
            "encoded_feature_names": self.feature_names(),
        }


class RidgeRegressor:
    def __init__(self, alpha: float):
        self.alpha = alpha
        self.weights: np.ndarray | None = None

    def fit(self, x: np.ndarray, y: np.ndarray) -> "RidgeRegressor":
        design = np.column_stack([np.ones(len(x)), x])
        penalty = np.eye(design.shape[1]) * self.alpha
        penalty[0, 0] = 0.0
        self.weights = np.linalg.solve(design.T @ design + penalty, design.T @ y)
        return self

    def predict(self, x: np.ndarray) -> np.ndarray:
        if self.weights is None:
            raise RuntimeError("RidgeRegressor has not been fitted.")
        design = np.column_stack([np.ones(len(x)), x])
        return design @ self.weights


class SoftmaxClassifier:
    """Multiclass logistic regression trained with full-batch Adam."""

    def __init__(self, learning_rate: float = 0.05, l2: float = 0.001):
        self.learning_rate = learning_rate
        self.l2 = l2
        self.weights: np.ndarray | None = None
        self.bias: np.ndarray | None = None
        self.epochs_trained = 0

    @staticmethod
    def _softmax(logits: np.ndarray) -> np.ndarray:
        shifted = logits - logits.max(axis=1, keepdims=True)
        exp = np.exp(shifted)
        return exp / exp.sum(axis=1, keepdims=True)

    def fit(
        self,
        x: np.ndarray,
        y: np.ndarray,
        x_validation: np.ndarray,
        y_validation: np.ndarray,
        epochs: int = 160,
        patience: int = 20,
    ) -> "SoftmaxClassifier":
        class_count = int(max(y.max(), y_validation.max())) + 1
        self.weights = np.zeros((x.shape[1], class_count), dtype=np.float64)
        self.bias = np.zeros(class_count, dtype=np.float64)
        m_w = np.zeros_like(self.weights)
        v_w = np.zeros_like(self.weights)
        m_b = np.zeros_like(self.bias)
        v_b = np.zeros_like(self.bias)
        best_loss = float("inf")
        best_parameters = None
        stale_epochs = 0

        for epoch in range(1, epochs + 1):
            probabilities = self._softmax(x @ self.weights + self.bias)
            probabilities[np.arange(len(y)), y] -= 1.0
            gradient_w = x.T @ probabilities / len(x) + self.l2 * self.weights
            gradient_b = probabilities.mean(axis=0)

            m_w = 0.9 * m_w + 0.1 * gradient_w
            v_w = 0.999 * v_w + 0.001 * gradient_w * gradient_w
            m_b = 0.9 * m_b + 0.1 * gradient_b
            v_b = 0.999 * v_b + 0.001 * gradient_b * gradient_b
            corrected_m_w = m_w / (1 - 0.9**epoch)
            corrected_v_w = v_w / (1 - 0.999**epoch)
            corrected_m_b = m_b / (1 - 0.9**epoch)
            corrected_v_b = v_b / (1 - 0.999**epoch)
            self.weights -= self.learning_rate * corrected_m_w / (np.sqrt(corrected_v_w) + 1e-8)
            self.bias -= self.learning_rate * corrected_m_b / (np.sqrt(corrected_v_b) + 1e-8)

            validation_probabilities = self.predict_proba(x_validation)
            clipped = np.clip(
                validation_probabilities[np.arange(len(y_validation)), y_validation],
                1e-12,
                1.0,
            )
            validation_loss = float(-np.log(clipped).mean())
            if validation_loss < best_loss - 1e-5:
                best_loss = validation_loss
                best_parameters = (self.weights.copy(), self.bias.copy(), epoch)
                stale_epochs = 0
            else:
                stale_epochs += 1
                if stale_epochs >= patience:
                    break

        if best_parameters is not None:
            self.weights, self.bias, self.epochs_trained = best_parameters
        return self

    def predict_proba(self, x: np.ndarray) -> np.ndarray:
        if self.weights is None or self.bias is None:
            raise RuntimeError("SoftmaxClassifier has not been fitted.")
        return self._softmax(x @ self.weights + self.bias)


def mae(actual: np.ndarray, predicted: np.ndarray) -> float:
    return float(np.mean(np.abs(actual - predicted)))


def accuracy(actual: np.ndarray, predicted: np.ndarray) -> float:
    return float(np.mean(actual == predicted))


def top_k_accuracy(actual: np.ndarray, probabilities: np.ndarray, k: int) -> float:
    top_k = np.argpartition(probabilities, -k, axis=1)[:, -k:]
    return float(np.mean(np.any(top_k == actual[:, None], axis=1)))


def save_model(
    path: Path,
    model_type: str,
    encoder: FeatureEncoder,
    weights: np.ndarray,
    bias: np.ndarray,
    target_classes: list[str] | None = None,
) -> None:
    metadata = {
        "model_type": model_type,
        "encoder": encoder.metadata(),
        "target_classes": target_classes,
    }
    np.savez_compressed(
        path,
        weights=weights,
        bias=bias,
        metadata_json=np.asarray(json.dumps(metadata)),
    )


def train(data_path: Path, output_dir: Path) -> dict:
    rows = load_rows(data_path)
    splits = {
        name: [row for row in rows if row["dataset_split"] == name]
        for name in ("train", "validation", "test")
    }
    if any(not split_rows for split_rows in splits.values()):
        raise ValueError("The workbook must contain train, validation, and test rows.")

    train_rows = splits["train"]
    validation_rows = splits["validation"]
    test_rows = splits["test"]
    encoder = FeatureEncoder(NUMERIC_FEATURES, CATEGORICAL_FEATURES).fit(train_rows)
    x_train = encoder.transform(train_rows)
    x_validation = encoder.transform(validation_rows)
    x_test = encoder.transform(test_rows)

    y_gap_train = np.asarray([row["target_next_gap_days"] for row in train_rows], dtype=float)
    y_gap_validation = np.asarray([row["target_next_gap_days"] for row in validation_rows], dtype=float)
    y_gap_test = np.asarray([row["target_next_gap_days"] for row in test_rows], dtype=float)

    ridge_candidates = []
    for alpha in (0.01, 0.1, 1.0, 10.0, 100.0):
        candidate = RidgeRegressor(alpha).fit(x_train, y_gap_train)
        ridge_candidates.append((mae(y_gap_validation, candidate.predict(x_validation)), alpha))
    _, selected_alpha = min(ridge_candidates)
    ridge = RidgeRegressor(selected_alpha).fit(x_train, y_gap_train)
    gap_predictions = np.clip(ridge.predict(x_test), 7, 180)
    personal_median_predictions = np.asarray(
        [row["median_gap_days"] for row in test_rows], dtype=float
    )

    product_classes = sorted({str(row["target_next_product_id"]) for row in train_rows})
    class_to_index = {product: index for index, product in enumerate(product_classes)}

    def encode_targets(target_rows: list[dict]) -> np.ndarray:
        unknown = {
            str(row["target_next_product_id"])
            for row in target_rows
            if str(row["target_next_product_id"]) not in class_to_index
        }
        if unknown:
            raise ValueError(f"Product targets absent from training data: {sorted(unknown)}")
        return np.asarray(
            [class_to_index[str(row["target_next_product_id"])] for row in target_rows],
            dtype=int,
        )

    y_product_train = encode_targets(train_rows)
    y_product_validation = encode_targets(validation_rows)
    y_product_test = encode_targets(test_rows)
    classifier = SoftmaxClassifier().fit(
        x_train,
        y_product_train,
        x_validation,
        y_product_validation,
    )
    product_probabilities = classifier.predict_proba(x_test)
    product_predictions = product_probabilities.argmax(axis=1)
    latest_product_predictions = np.asarray(
        [class_to_index[str(row["latest_product_id"])] for row in test_rows], dtype=int
    )
    most_common_class = int(np.bincount(y_product_train).argmax())
    most_common_predictions = np.full(len(test_rows), most_common_class)

    output_dir.mkdir(parents=True, exist_ok=True)
    if ridge.weights is None or classifier.weights is None or classifier.bias is None:
        raise RuntimeError("Training did not produce model parameters.")
    save_model(
        output_dir / "purchase_timing_ridge.npz",
        "ridge_regression",
        encoder,
        ridge.weights[1:],
        np.asarray([ridge.weights[0]]),
    )
    save_model(
        output_dir / "next_product_softmax.npz",
        "multiclass_logistic_regression",
        encoder,
        classifier.weights,
        classifier.bias,
        product_classes,
    )

    metrics = {
        "data_source": str(data_path),
        "split_strategy": "predefined customer-level train/validation/test split",
        "rows": {name: len(split_rows) for name, split_rows in splits.items()},
        "customers": {
            name: len({row["customer_id"] for row in split_rows})
            for name, split_rows in splits.items()
        },
        "purchase_timing": {
            "target": "target_next_gap_days",
            "metric": "mean_absolute_error_days",
            "personal_median_baseline_test": round(
                mae(y_gap_test, personal_median_predictions), 3
            ),
            "ridge_test": round(mae(y_gap_test, gap_predictions), 3),
            "selected_alpha": selected_alpha,
        },
        "next_product": {
            "target": "target_next_product_id",
            "classes": len(product_classes),
            "most_common_baseline_test_accuracy": round(
                accuracy(y_product_test, most_common_predictions), 4
            ),
            "repeat_latest_baseline_test_accuracy": round(
                accuracy(y_product_test, latest_product_predictions), 4
            ),
            "softmax_test_accuracy": round(
                accuracy(y_product_test, product_predictions), 4
            ),
            "softmax_test_top3_accuracy": round(
                top_k_accuracy(y_product_test, product_probabilities, 3), 4
            ),
            "epochs_trained": classifier.epochs_trained,
        },
    }
    (output_dir / "metrics.json").write_text(
        json.dumps(metrics, indent=2), encoding="utf-8"
    )

    with (output_dir / "test_predictions.csv").open(
        "w", newline="", encoding="utf-8"
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "example_id",
                "customer_id",
                "actual_gap_days",
                "predicted_gap_days",
                "actual_product_id",
                "predicted_product_id",
                "predicted_product_confidence",
            ],
        )
        writer.writeheader()
        for index, row in enumerate(test_rows):
            predicted_class = int(product_predictions[index])
            writer.writerow(
                {
                    "example_id": row["example_id"],
                    "customer_id": row["customer_id"],
                    "actual_gap_days": int(y_gap_test[index]),
                    "predicted_gap_days": round(float(gap_predictions[index]), 2),
                    "actual_product_id": product_classes[int(y_product_test[index])],
                    "predicted_product_id": product_classes[predicted_class],
                    "predicted_product_confidence": round(
                        float(product_probabilities[index, predicted_class]), 4
                    ),
                }
            )
    return metrics


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", type=Path, default=DEFAULT_DATA)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    arguments = parse_args()
    metrics = train(arguments.data.resolve(), arguments.output.resolve())
    print(json.dumps(metrics, indent=2))


if __name__ == "__main__":
    main()
