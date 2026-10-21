"""Score current customers and prepare a review-only pre-MantaGO campaign.

No email or MantaGO request is sent. The script produces auditable JSON/CSV
files that a later website and MantaGO integration can consume.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from statistics import mean, median, pstdev
from urllib.parse import urlencode

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DATA = ROOT / "data" / "retail_purchase_training_dataset.xlsx"
DEFAULT_ARTIFACTS = ROOT / "prediction" / "artifacts"
DEFAULT_CONFIG = ROOT / "prediction" / "policy_config.json"
DEFAULT_HISTORY = ROOT / "prediction" / "campaign_history.json"
DEFAULT_OUTPUT = ROOT / "prediction" / "campaign_output"


def load_sheet(workbook, sheet_name: str) -> list[dict]:
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    return [dict(zip(headers, values)) for values in iterator]


def load_model(path: Path) -> dict:
    archive = np.load(path)
    return {
        "weights": archive["weights"],
        "bias": archive["bias"],
        "metadata": json.loads(str(archive["metadata_json"])),
    }


def encode_features(rows: list[dict], encoder: dict) -> np.ndarray:
    numeric_names = encoder["numeric_features"]
    numeric = np.asarray(
        [
            [float(bool(row[name])) if isinstance(row[name], bool) else float(row[name])
             for name in numeric_names]
            for row in rows
        ],
        dtype=np.float64,
    )
    means = np.asarray(encoder["means"], dtype=np.float64)
    scales = np.asarray(encoder["scales"], dtype=np.float64)
    blocks = [(numeric - means) / scales]
    for feature in encoder["categorical_features"]:
        categories = encoder["categories"][feature]
        indexes = {category: index for index, category in enumerate(categories)}
        encoded = np.zeros((len(rows), len(categories)), dtype=np.float64)
        for row_index, row in enumerate(rows):
            category_index = indexes.get(str(row[feature]))
            if category_index is not None:
                encoded[row_index, category_index] = 1.0
        blocks.append(encoded)
    return np.column_stack(blocks)


def softmax(logits: np.ndarray) -> np.ndarray:
    shifted = logits - logits.max(axis=1, keepdims=True)
    exponentials = np.exp(shifted)
    return exponentials / exponentials.sum(axis=1, keepdims=True)


def build_current_features(
    customers: list[dict], orders: list[dict], as_of: datetime
) -> list[dict]:
    grouped_orders: dict[str, list[dict]] = defaultdict(list)
    for order in orders:
        grouped_orders[str(order["customer_id"])].append(order)
    for history in grouped_orders.values():
        history.sort(key=lambda order: order["order_date"])

    rows = []
    for customer in customers:
        customer_id = str(customer["customer_id"])
        history = grouped_orders[customer_id]
        if not history:
            continue
        gaps = [
            (history[index]["order_date"] - history[index - 1]["order_date"]).days
            for index in range(1, len(history))
        ]
        product_counts = Counter(str(order["product_id"]) for order in history)
        dominant_product_id, dominant_count = product_counts.most_common(1)[0]
        latest = history[-1]
        previous = history[-2] if len(history) >= 2 else latest
        rows.append(
            {
                "customer_id": customer_id,
                "customer_name": str(customer["customer_name"]),
                "dataset_split": str(customer["dataset_split"]),
                "as_of_date": as_of,
                "latest_order_date": latest["order_date"],
                "days_since_last_purchase": (as_of - latest["order_date"]).days,
                "model_ready": len(history) >= 3,
                "order_count_to_date": len(history),
                "days_since_previous_order": gaps[-1] if gaps else 0,
                "mean_gap_days": round(mean(gaps), 2) if gaps else 0.0,
                "median_gap_days": round(median(gaps), 2) if gaps else 0.0,
                "gap_std_days": round(pstdev(gaps), 2) if gaps else 0.0,
                "dominant_product_id": dominant_product_id,
                "dominant_product_share": round(dominant_count / len(history), 4),
                "latest_product_id": str(latest["product_id"]),
                "previous_product_id": str(previous["product_id"]),
                "same_as_previous_product": latest["product_id"] == previous["product_id"],
                "latest_quantity": int(latest["quantity"]),
                "latest_order_value_sgd": float(latest["order_value_sgd"]),
                "latest_discount_pct": float(latest["discount_pct"]),
                "latest_channel": str(latest["channel"]),
                "region": str(customer["region"]),
                "customer_profile": str(customer["customer_profile"]),
                "contact_consent": bool(customer["contact_consent"]),
            }
        )
    return rows


def stable_holdout(customer_id: str, percentage: int) -> bool:
    bucket = int(hashlib.sha256(customer_id.encode("utf-8")).hexdigest()[:8], 16) % 100
    return bucket < percentage


def policy_decision(row: dict, config: dict) -> tuple[str, str, str]:
    if not row["contact_consent"]:
        return "suppress", "no_contact_consent", "Customer has no marketing consent."
    if row["order_count_to_date"] < config["minimum_orders"]:
        return "stay_quiet", "insufficient_history", "Too few orders for reliable timing."
    if not row["has_available_recommendation"]:
        return "stay_quiet", "no_available_product", "No approved predicted product is currently in stock."
    if (
        row["days_since_last_campaign"] is not None
        and row["days_since_last_campaign"] < config["contact_cooldown_days"]
    ):
        return "stay_quiet", "contact_cooldown", "Customer was contacted too recently."

    days_until_due = row["days_until_predicted_purchase"]
    if -config["contact_window_after_days"] <= days_until_due <= config["contact_window_before_days"]:
        return "eligible", "due_soon", "Customer is inside the approved purchase window."
    if days_until_due < -config["lapsed_after_days"]:
        return "eligible", "lapsed", "Customer is beyond the predicted purchase cycle."
    return "stay_quiet", "not_due", "Customer is not currently due for contact."


def opportunity_score(row: dict) -> float:
    state_weight = 1.0 if row["customer_state"] == "lapsed" else 0.8
    confidence = float(row["product_confidence"])
    product_stability = float(row["dominant_product_share"])
    return round(0.5 * state_weight + 0.3 * confidence + 0.2 * product_stability, 4)


def demo_email(customer_id: str) -> str:
    return f"{customer_id.lower()}@demo.example"


def json_ready(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, dict):
        return {key: json_ready(item) for key, item in value.items()}
    if isinstance(value, list):
        return [json_ready(item) for item in value]
    return value


def commerce_context(product_id: str | None, products: dict, inventory: dict, offers: dict):
    if product_id is None:
        return None
    product = products[product_id]
    stock = inventory[product_id]
    offer = offers.get(product_id)
    return {
        "product_id": product_id,
        "product_name": str(product["product_name"]),
        "category": str(product["category"]),
        "stock_quantity": int(stock["stock_quantity"]),
        "stock_status": str(stock["stock_status"]),
        "next_restock_date": stock["next_restock_date"],
        "regular_price_sgd": round(float(product["unit_price_sgd"]), 2),
        "current_price_sgd": round(
            float(offer["discount_price_sgd"]) if offer else float(product["unit_price_sgd"]), 2
        ),
        "discount_pct": round(float(offer["discount_pct"]), 4) if offer else 0.0,
        "offer_id": str(offer["offer_id"]) if offer else None,
        "offer_valid_until": offer["end_date"] if offer else None,
        "offer_terms": str(offer["terms"]) if offer else None,
        "inventory_last_updated": stock["last_updated"],
        "offer_last_updated": offer["last_updated"] if offer else None,
    }


def write_json(path: Path, value) -> None:
    path.write_text(json.dumps(json_ready(value), indent=2), encoding="utf-8")


def score(
    data_path: Path,
    artifact_dir: Path,
    config_path: Path,
    history_path: Path,
    output_dir: Path,
) -> dict:
    config = json.loads(config_path.read_text(encoding="utf-8"))
    campaign_history = json.loads(history_path.read_text(encoding="utf-8"))
    workbook = load_workbook(data_path, read_only=True, data_only=True)
    customers = load_sheet(workbook, "Customers")
    orders = load_sheet(workbook, "Raw Orders")
    products = load_sheet(workbook, "Products")
    inventory_rows = load_sheet(workbook, "Inventory")
    offer_rows = load_sheet(workbook, "Active Offers")
    snapshots = load_sheet(workbook, "Current Snapshot")
    workbook.close()

    as_of = max(snapshot["as_of_date"] for snapshot in snapshots)
    last_campaign_by_customer = {}
    for event in campaign_history:
        if event.get("status") != "sent":
            continue
        sent_at = datetime.fromisoformat(event["sent_at"].replace("Z", "+00:00"))
        sent_at = sent_at.replace(tzinfo=None)
        customer_id = str(event["customer_id"])
        current = last_campaign_by_customer.get(customer_id)
        if current is None or sent_at > current:
            last_campaign_by_customer[customer_id] = sent_at
    features = build_current_features(customers, orders, as_of)
    product_catalog = {str(product["product_id"]): product for product in products}
    inventory_catalog = {str(row["product_id"]): row for row in inventory_rows}
    active_offers = {
        str(row["product_id"]): row
        for row in offer_rows
        if bool(row["is_active"]) and row["start_date"] <= as_of <= row["end_date"]
    }

    def available(product_id: str) -> bool:
        stock = inventory_catalog.get(product_id)
        return bool(stock) and int(stock["stock_quantity"]) > 0 and stock["stock_status"] != "out_of_stock"

    timing_model = load_model(artifact_dir / "purchase_timing_ridge.npz")
    product_model = load_model(artifact_dir / "next_product_softmax.npz")
    encoder = product_model["metadata"]["encoder"]
    ready_features = [row for row in features if row["model_ready"]]
    encoded = encode_features(ready_features, encoder)

    predicted_gaps = np.clip(
        encoded @ timing_model["weights"] + float(timing_model["bias"][0]), 7, 180
    )
    probabilities = softmax(encoded @ product_model["weights"] + product_model["bias"])
    product_classes = product_model["metadata"]["target_classes"]

    model_results = {}
    for index, feature_row in enumerate(ready_features):
        ranked_indexes = np.argsort(probabilities[index])[::-1][:3]
        top_products = [
            {
                "product_id": product_classes[int(class_index)],
                "product_name": str(product_catalog[product_classes[int(class_index)]]["product_name"]),
                "probability": round(float(probabilities[index, class_index]), 4),
            }
            for class_index in ranked_indexes
        ]
        predicted_product = top_products[0]
        low_confidence = predicted_product["probability"] < config["minimum_product_confidence"]
        ranked_product_ids = [product["product_id"] for product in top_products]
        candidate_product_ids = (
            [feature_row["latest_product_id"], feature_row["dominant_product_id"], *ranked_product_ids]
            if low_confidence else ranked_product_ids
        )
        safe_product_id = next((product_id for product_id in candidate_product_ids if available(product_id)), None)
        safe_product = product_catalog[safe_product_id] if safe_product_id else None
        approved_upsell_id = str(safe_product["approved_recommendation_id"]) if safe_product else None
        upsell_product_id = approved_upsell_id if approved_upsell_id and available(approved_upsell_id) else None
        predicted_gap = round(float(predicted_gaps[index]), 2)
        days_until_due = round(predicted_gap - feature_row["days_since_last_purchase"], 2)
        model_results[feature_row["customer_id"]] = {
            "predicted_gap_days": predicted_gap,
            "days_until_predicted_purchase": days_until_due,
            "predicted_next_purchase_date": (
                feature_row["latest_order_date"]
                + timedelta(days=round(predicted_gap))
            ),
            "predicted_product_id": predicted_product["product_id"],
            "predicted_product_name": predicted_product["product_name"],
            "product_confidence": predicted_product["probability"],
            "top_3_products": top_products,
            "recommendation_strategy": (
                "no_available_product" if safe_product_id is None
                else "available_product_fallback" if safe_product_id != predicted_product["product_id"]
                else "model"
            ),
            "has_available_recommendation": safe_product_id is not None,
            "safe_reorder_product_id": safe_product_id,
            "safe_reorder_product_name": str(safe_product["product_name"]) if safe_product else None,
            "approved_upsell_product_id": upsell_product_id,
            "approved_upsell_product_name": str(product_catalog[upsell_product_id]["product_name"]) if upsell_product_id else None,
            "reorder_commerce": commerce_context(safe_product_id, product_catalog, inventory_catalog, active_offers),
            "upsell_commerce": commerce_context(upsell_product_id, product_catalog, inventory_catalog, active_offers),
        }

    scored = []
    for feature_row in features:
        last_campaign = last_campaign_by_customer.get(feature_row["customer_id"])
        feature_row["last_campaign_contact_date"] = last_campaign
        feature_row["days_since_last_campaign"] = (
            (as_of - last_campaign).days if last_campaign is not None else None
        )
        if feature_row["model_ready"]:
            prediction = model_results[feature_row["customer_id"]]
        else:
            latest_product = product_catalog[feature_row["latest_product_id"]]
            upsell_product_id = str(latest_product["approved_recommendation_id"])
            prediction = {
                "predicted_gap_days": None,
                "days_until_predicted_purchase": None,
                "predicted_next_purchase_date": None,
                "predicted_product_id": None,
                "predicted_product_name": None,
                "product_confidence": 0.0,
                "top_3_products": [],
                "recommendation_strategy": "insufficient_history",
                "has_available_recommendation": available(feature_row["latest_product_id"]),
                "safe_reorder_product_id": feature_row["latest_product_id"],
                "safe_reorder_product_name": str(latest_product["product_name"]),
                "approved_upsell_product_id": upsell_product_id if available(upsell_product_id) else None,
                "approved_upsell_product_name": str(product_catalog[upsell_product_id]["product_name"]) if available(upsell_product_id) else None,
                "reorder_commerce": commerce_context(feature_row["latest_product_id"], product_catalog, inventory_catalog, active_offers) if available(feature_row["latest_product_id"]) else None,
                "upsell_commerce": commerce_context(upsell_product_id, product_catalog, inventory_catalog, active_offers) if available(upsell_product_id) else None,
            }
        row = {**feature_row, **prediction}
        status, state, reason = policy_decision(row, config)
        row.update(
            {
                "policy_status": status,
                "customer_state": state,
                "policy_reason": reason,
                "holdout": False,
                "opportunity_score": 0.0,
            }
        )
        if status == "eligible":
            row["opportunity_score"] = opportunity_score(row)
            if stable_holdout(row["customer_id"], config["holdout_percentage"]):
                row["policy_status"] = "holdout"
                row["holdout"] = True
                row["policy_reason"] = "Eligible customer assigned to the measurement holdout."
        scored.append(row)

    send_candidates = sorted(
        (row for row in scored if row["policy_status"] == "eligible"),
        key=lambda row: (-row["opportunity_score"], row["customer_id"]),
    )
    selected_ids = {
        row["customer_id"]
        for row in send_candidates[: config["maximum_campaign_recipients"]]
    }
    for row in send_candidates[config["maximum_campaign_recipients"] :]:
        row["policy_status"] = "stay_quiet"
        row["customer_state"] = "campaign_capacity"
        row["policy_reason"] = "Eligible but outside this campaign's configured capacity."

    selected = [row for row in scored if row["customer_id"] in selected_ids]
    not_selected = [row for row in scored if row["customer_id"] not in selected_ids]
    campaign_id = f"RET-{as_of:%Y%m%d}"
    generated_at = datetime.now(timezone.utc).isoformat()
    email_queue = []
    landing_context = {}
    mantago_customers = []
    for row in selected:
        token = hashlib.sha256(
            f"{campaign_id}:{row['customer_id']}:synthetic-demo-only".encode("utf-8")
        ).hexdigest()[:32]
        landing_url = f"{config['landing_page_base_url']}?{urlencode({'t': token})}"
        first_name = row["customer_name"].split()[0]
        conversation_goal = (
            "retention_recovery" if row["customer_state"] == "lapsed" else "replenishment_upsell"
        )
        subject = (
            f"Need help restocking {row['safe_reorder_product_name']}?"
            if row["customer_state"] == "lapsed"
            else f"Your {row['safe_reorder_product_name']} may be due soon"
        )
        email_queue.append(
            {
                "campaign_id": campaign_id,
                "customer_id": row["customer_id"],
                "recipient_email": demo_email(row["customer_id"]),
                "recipient_is_synthetic": True,
                "status": "pending_review",
                "subject": subject,
                "plain_text": (
                    f"Hi {first_name}, based on your usual purchase cycle, your "
                    f"{row['safe_reorder_product_name']} may be due. Review your recommendation: "
                    f"{landing_url}"
                ),
                "landing_url": landing_url,
                "conversation_goal": conversation_goal,
            }
        )
        landing_context[token] = {
            "campaign_id": campaign_id,
            "customer_id": row["customer_id"],
            "customer_first_name": first_name,
            "conversation_goal": conversation_goal,
            "customer_state": row["customer_state"],
            "reorder_product_id": row["safe_reorder_product_id"],
            "reorder_product_name": row["safe_reorder_product_name"],
            "upsell_product_id": row["approved_upsell_product_id"],
            "upsell_product_name": row["approved_upsell_product_name"],
            "reorder_product": row["reorder_commerce"],
            "upsell_product": row["upsell_commerce"],
            "product_confidence": row["product_confidence"],
            "human_handover_available": True,
            "data_sources": ["Products", "Inventory", "Active Offers"],
        }
        mantago_customers.append({
            "customer_id": row["customer_id"],
            "mantago_client_id": None,
            "conversation_goal": conversation_goal,
            "customer_state": row["customer_state"],
            "reorder_product": row["reorder_commerce"],
            "upsell_product": row["upsell_commerce"],
            "human_handover_available": True,
        })

    output_dir.mkdir(parents=True, exist_ok=True)
    write_json(output_dir / "scored_customers.json", scored)
    write_json(output_dir / "eligible_customers.json", selected)
    write_json(output_dir / "suppressed_customers.json", not_selected)
    write_json(output_dir / "email_campaign_queue.json", email_queue)
    write_json(output_dir / "landing_context.json", landing_context)
    write_json(output_dir / "mantago_ready_payload.json", {
        "schema_version": "1.0",
        "campaign_id": campaign_id,
        "mode": "not_sent_waiting_for_mantago_access",
        "customers": mantago_customers,
    })

    with (output_dir / "email_campaign_queue.csv").open(
        "w", newline="", encoding="utf-8"
    ) as handle:
        fieldnames = [
            "campaign_id", "customer_id", "recipient_email", "recipient_is_synthetic",
            "status", "subject", "plain_text", "landing_url", "conversation_goal",
        ]
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(email_queue)

    status_counts = Counter(row["policy_status"] for row in scored)
    state_counts = Counter(row["customer_state"] for row in scored)
    summary = {
        "campaign_id": campaign_id,
        "generated_at": generated_at,
        "mode": "review_only_no_messages_sent",
        "customers_scored": len(scored),
        "campaign_history_events_checked": len(campaign_history),
        "selected_for_email_review": len(selected),
        "status_counts": dict(sorted(status_counts.items())),
        "state_counts": dict(sorted(state_counts.items())),
        "policy_config": config,
        "outputs": [
            "scored_customers.json",
            "eligible_customers.json",
            "suppressed_customers.json",
            "email_campaign_queue.json",
            "email_campaign_queue.csv",
            "landing_context.json",
            "mantago_ready_payload.json",
        ],
    }
    write_json(output_dir / "campaign_summary.json", summary)
    return summary


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", type=Path, default=DEFAULT_DATA)
    parser.add_argument("--artifacts", type=Path, default=DEFAULT_ARTIFACTS)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--history", type=Path, default=DEFAULT_HISTORY)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    arguments = parse_args()
    summary = score(
        arguments.data.resolve(),
        arguments.artifacts.resolve(),
        arguments.config.resolve(),
        arguments.history.resolve(),
        arguments.output.resolve(),
    )
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
